import io
import re
from datetime import datetime, timezone
from typing import List, Dict, Any
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    """Sets background shading of a docx table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def generate_proposal_docx(
    proposal_title: str,
    client_name: str,
    project_title: str,
    status: str,
    sections: List[Dict[str, Any]],
    branding: Dict[str, Any] = None
) -> bytes:
    """Generates an editable, professionally formatted Word (.docx) document."""
    doc = docx.Document()

    # Set 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("ProposalCraft AI — Statement of Work")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Commercial Confidential — SOW Draft")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(100, 116, 139)

    # Document Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(proposal_title)
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.color.rgb = RGBColor(15, 36, 56) # Deep Navy
    p_title.paragraph_format.space_after = Pt(4)

    # Subtitle / Metadata
    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run(f"Prepared for: {client_name}  |  Project: {project_title}  |  Date: {datetime.now(timezone.utc).strftime('%B %d, %Y')}")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    p_sub.paragraph_format.space_after = Pt(16)

    # Render each proposal section
    for sec in sorted(sections, key=lambda s: s.get('sort_order', 0)):
        sec_title = sec.get('title', 'Section')
        content_md = sec.get('content_markdown', '')

        # Heading 1
        h1 = doc.add_heading(level=1)
        r_h1 = h1.add_run(sec_title)
        r_h1.font.name = "Calibri"
        r_h1.font.size = Pt(14)
        r_h1.font.color.rgb = RGBColor(43, 102, 149) # Muted Blue
        h1.paragraph_format.space_before = Pt(14)
        h1.paragraph_format.space_after = Pt(6)

        lines = content_md.split('\n')
        in_table = False
        table_rows = []

        for line in lines:
            trimmed = line.strip()
            if not trimmed:
                if in_table and table_rows:
                    render_docx_table(doc, table_rows)
                    table_rows = []
                    in_table = False
                continue

            # Check table lines
            if trimmed.startswith('|') and trimmed.endswith('|'):
                if re.match(r'^\|[\s\-:]+\|\s*$', trimmed):
                    continue
                in_table = True
                cells = [c.strip() for c in trimmed[1:-1].split('|')]
                table_rows.append(cells)
                continue
            else:
                if in_table and table_rows:
                    render_docx_table(doc, table_rows)
                    table_rows = []
                    in_table = False

            # Check Headings
            if trimmed.startswith('### '):
                h2 = doc.add_heading(level=2)
                r_h2 = h2.add_run(trimmed[4:])
                r_h2.font.name = "Calibri"
                r_h2.font.size = Pt(12)
                r_h2.font.color.rgb = RGBColor(15, 36, 56)
                h2.paragraph_format.space_before = Pt(10)
                h2.paragraph_format.space_after = Pt(4)
            elif trimmed.startswith('#### '):
                p = doc.add_paragraph()
                r = p.add_run(trimmed[5:])
                r.bold = True
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)
                p.paragraph_format.space_after = Pt(3)
            elif trimmed.startswith('- ') or trimmed.startswith('* '):
                p = doc.add_paragraph(style='List Bullet')
                cleaned = re.sub(r'\*\*(.*?)\*\*', r'\1', trimmed[2:])
                r = p.add_run(cleaned)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                p.paragraph_format.space_after = Pt(2)
            elif trimmed.startswith('> '):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                cleaned = re.sub(r'\*\*(.*?)\*\*', r'\1', trimmed[2:])
                r = p.add_run(cleaned)
                r.italic = True
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(30, 64, 175)
                p.paragraph_format.space_after = Pt(4)
            else:
                p = doc.add_paragraph()
                cleaned = re.sub(r'\*\*(.*?)\*\*', r'\1', trimmed)
                r = p.add_run(cleaned)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                p.paragraph_format.space_after = Pt(4)

        if in_table and table_rows:
            render_docx_table(doc, table_rows)
            table_rows = []

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def render_docx_table(doc: docx.Document, rows: List[List[str]]):
    """Creates a styled Word table."""
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for row_idx, row in enumerate(rows):
        for col_idx, cell_value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = cell_value
            p = cell.paragraphs[0]
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(9)
            if row_idx == 0:
                p.runs[0].bold = True
                set_cell_background(cell, "F1F5F9")
            else:
                set_cell_background(cell, "FFFFFF" if row_idx % 2 == 1 else "F8FAFC")

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)
